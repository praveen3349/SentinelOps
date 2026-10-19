from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output"
DOCX_PATH = OUTPUT / "Final_SentinelOps_IEEE_Paper.docx"
FIGURE_PATH = OUTPUT / "sentinelops_delivery_architecture.png"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    tc_pr.append(shade)


def set_cell_border(cell, color="B7C3D0"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def set_table_cell_text(cell, text, bold=False, color=None, size=7.6):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_column_count(section, count=2, spacing=0.22):
    sect_pr = section._sectPr
    cols = sect_pr.xpath("./w:cols")
    cols_el = cols[0] if cols else OxmlElement("w:cols")
    cols_el.set(qn("w:num"), str(count))
    cols_el.set(qn("w:space"), str(int(spacing * 1440)))
    if not cols:
        sect_pr.append(cols_el)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_text(doc, text, size=9.5, italic=False, bold=False, alignment=None, after=3, before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.0
    if alignment is not None:
        p.alignment = alignment
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.italic = italic
    run.bold = bold
    return p


def add_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(9.4)
    return p


def add_label_paragraph(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    label_run = p.add_run(label)
    label_run.bold = True
    label_run.font.name = "Times New Roman"
    label_run.font.size = Pt(9.2)
    body_run = p.add_run(text)
    body_run.font.name = "Times New Roman"
    body_run.font.size = Pt(9.2)
    for run in (label_run, body_run):
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    return p


def add_architecture_flow(doc):
    caption = add_text(doc, "Fig. 1. SentinelOps controlled delivery flow.", size=7.6, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER, after=2, before=3)
    caption.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    values = ("Repository intake\nand stack analysis", "Local validation\nand service URLs", "Delivery assets\nand operations plan")
    colors = ("EAF1F8", "E8F5F1", "F0EBF8")
    for col_idx, cell in enumerate(table.rows[0].cells):
        set_cell_border(cell, "9BB5C8")
        set_cell_shading(cell, colors[col_idx])
        set_table_cell_text(cell, values[col_idx], bold=True, color="17324D", size=7.5)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    return table


def add_module_table(doc):
    caption = add_text(doc, "TABLE I.  SENTINELOPS MODULES AND GENERATED OUTPUTS", size=7.8, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER, after=2, before=3)
    caption.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.18)
    table.columns[1].width = Inches(1.83)
    headers = [("Module", 0), ("Output or decision", 1)]
    for text, idx in headers:
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, "1F4E79")
        set_cell_border(cell, "1F4E79")
        set_table_cell_text(cell, text, bold=True, color="FFFFFF", size=7.4)
    rows = [
        ("Repository intake", "Cloned workspace, commit metadata, and project inventory"),
        ("Stack analysis", "Editable language, framework, package, and frontend/backend findings"),
        ("Verification", "Install and start commands, health result, and local access URLs"),
        ("Delivery assets", "Dockerfile, Compose file, CI workflow, Jenkinsfile, or GitLab pipeline"),
        ("Operations", "Deployment profile, HPA manifest, Prometheus configuration, dashboard JSON, and runbook"),
    ]
    for left, right in rows:
        cells = table.add_row().cells
        for cell in cells:
            set_cell_border(cell)
        set_cell_shading(cells[0], "EAF1F8")
        set_table_cell_text(cells[0], left, bold=True)
        set_table_cell_text(cells[1], right)
    for row in table.rows:
        row.cells[0].width = Inches(1.18)
        row.cells[1].width = Inches(1.83)


def add_validation_table(doc):
    caption = add_text(doc, "TABLE II  PROPOSED REPOSITORY EVALUATION CASES", size=7.8, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER, after=2, before=3)
    caption.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ("Project shape", "Expected finding", "Evidence reviewed")
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, "1F4E79")
        set_cell_border(cell, "1F4E79")
        set_table_cell_text(cell, text, bold=True, color="FFFFFF", size=7.1)
    rows = [
        ("Single API", "Runtime and port", "Manifest, scripts, source entry point"),
        ("Web client", "Framework and build command", "Dependencies, lock file, public assets"),
        ("Client plus server", "Two services and URLs", "Independent manifests and folders"),
        ("Container-ready repo", "Existing files preserved", "Docker and Compose configuration"),
    ]
    for values in rows:
        cells = table.add_row().cells
        for cell in cells:
            set_cell_border(cell)
        for index, value in enumerate(values):
            if index == 0:
                set_cell_shading(cells[index], "EAF1F8")
            set_table_cell_text(cells[index], value, bold=(index == 0), size=7.0)


def add_profile_table(doc):
    caption = add_text(doc, "TABLE III  DEPLOYMENT PROFILES EXPOSED BY SENTINELOPS", size=7.8, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER, after=2, before=3)
    caption.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ("Profile", "Appropriate use", "Generated guidance")
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, "1F4E79")
        set_cell_border(cell, "1F4E79")
        set_table_cell_text(cell, text, bold=True, color="FFFFFF", size=7.1)
    rows = [
        ("Native local", "Early debugging on one workstation", "Commands, URLs, health status, stop action"),
        ("Docker Compose", "Repeatable local service bundle", "Dockerfile, Compose file, image and port mapping"),
        ("Kubernetes", "Cluster deployment with approval", "Deployment, Service, HPA, capacity policy"),
    ]
    for values in rows:
        cells = table.add_row().cells
        for cell in cells:
            set_cell_border(cell)
        for index, value in enumerate(values):
            if index == 0:
                set_cell_shading(cells[index], "EAF1F8")
            set_table_cell_text(cells[index], value, bold=(index == 0), size=7.0)


def add_references(doc):
    references = [
        '[1] K. Raj, K. P. Poojary, Abhay, A. R. S., and L. K. S. R., “AI Driven DevOps for Intelligent Automation in Continuous Software Delivery Pipelines,” in Proc. 3rd Int. Conf. on Emerging Computing Systems and Network, 2025.',
        '[2] J. Humble and D. Farley, Continuous Delivery Reliable Software Releases through Build Test and Deployment Automation. Boston, MA, USA: Addison-Wesley, 2010.',
        '[3] D. Merkel, “Docker Lightweight Linux Containers for Consistent Development and Deployment,” Linux Journal, vol. 2014, no. 239, 2014.',
        '[4] Prometheus Authors, “Overview,” Prometheus Documentation. [Online]. Available: https://prometheus.io/docs/introduction/overview/. Accessed: Sep. 24, 2026.',
        '[5] The Kubernetes Authors, “Horizontal Pod Autoscaling,” Kubernetes Documentation. [Online]. Available: https://kubernetes.io/docs/concepts/workloads/autoscaling/horizontal-pod-autoscale/. Accessed: Sep. 24, 2026.',
        '[6] Grafana Labs, “Build Your First Dashboard,” Grafana Documentation. [Online]. Available: https://grafana.com/docs/grafana/latest/fundamentals/getting-started/first-dashboards/. Accessed: Sep. 24, 2026.',
    ]
    for reference in references:
        p = add_text(doc, reference, size=7.7, after=2)
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.first_line_indent = Inches(-0.15)


def build_document():
    OUTPUT.mkdir(exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.54)
    section.bottom_margin = Inches(0.53)
    section.left_margin = Inches(0.58)
    section.right_margin = Inches(0.58)
    section.header_distance = Inches(0.24)
    section.footer_distance = Inches(0.26)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(8.7)
    normal.paragraph_format.space_after = Pt(0)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("SentinelOps IEEE Style Project Paper  ")
    footer_run.font.name = "Times New Roman"
    footer_run.font.size = Pt(7.3)
    add_page_number(footer)

    title = add_text(doc, "SentinelOps An AI Assisted Desktop Control Plane for Automated DevOps Delivery", size=17.5, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER, after=7)
    title.paragraph_format.line_spacing = 0.94
    for line in [
        "Subin and Praveen",
        "Sathyabama Institute of Science and Technology Deemed to be University",
        "Chennai India",
    ]:
        p = add_text(doc, line, size=9.2, alignment=WD_ALIGN_PARAGRAPH.CENTER, after=1)
        p.paragraph_format.line_spacing = 0.92

    body_section = doc.add_section(WD_SECTION.CONTINUOUS)
    body_section.top_margin = Inches(0.54)
    body_section.bottom_margin = Inches(0.53)
    body_section.left_margin = Inches(0.58)
    body_section.right_margin = Inches(0.58)
    set_column_count(body_section, 2)
    footer = body_section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if not footer.text:
        footer_run = footer.add_run("SentinelOps IEEE Style Project Paper  ")
        footer_run.font.name = "Times New Roman"
        footer_run.font.size = Pt(7.3)
        add_page_number(footer)

    add_label_paragraph(doc, "Abstract—", "SentinelOps is a desktop prototype that turns an unfamiliar Git repository into a reviewed delivery plan. It clones a selected repository, identifies languages, frameworks, package managers, service boundaries, and likely start commands, then presents those findings for correction. The system validates the application locally before it generates delivery assets. Depending on the selected profile, SentinelOps prepares Docker and Compose files, CI definitions for GitHub Actions, GitLab CI, or Jenkins, Kubernetes manifests with an autoscaling policy, and observability material for Prometheus and Grafana. The design keeps installation and execution under user control: Docker remains an external prerequisite on Windows, and the application records commands, URLs, and health checks instead of hiding them. This paper describes the architecture, the module boundaries, and an evaluation plan for repository analysis and deployment readiness. The contribution is a practical, repo-aware control plane that makes the transition from source code to an inspectable DevOps workflow easier for student and small-team projects. It does not claim unattended production remediation or cloud deployment without explicit credentials and approval.")
    add_label_paragraph(doc, "Index Terms—", "DevOps automation, repository analysis, CI/CD, containerization, observability, desktop application.")

    add_heading(doc, "I INTRODUCTION")
    add_text(doc, "Many small software teams can write and run an application locally, yet still struggle to package it, choose a deployment target, create a continuous integration pipeline, and observe the running service. Those tasks require information that is scattered across package files, scripts, environment files, and repository layout. A developer who has not seen the project before may also miss that a React client and an API must start independently. Continuous delivery practices address this gap by treating build, test, and deployment work as part of software delivery rather than as a final manual step [2].")
    add_text(doc, "The reference work in [1] motivates the use of AI in DevOps decisions. SentinelOps narrows that idea to a desktop workflow that a user can inspect. It uses deterministic repository evidence first, then exposes an editable model of the detected stack. This choice is deliberate. A useful automation tool should show why it inferred a technology or command and should let the project owner correct it before it creates delivery files or starts containers.")

    add_heading(doc, "II RELATED WORK AND MOTIVATION")
    add_text(doc, "Continuous delivery is commonly described as the ability to keep software in a releasable state through an automated build, test, and deployment process [2]. In practice, smaller projects often begin with only a source repository and a local start command. The gap between that starting point and a repeatable delivery process is larger than it appears. It includes dependency installation, environment configuration, packaging, test selection, image publishing, deployment selection, and operational visibility. Each step depends on details that may be implicit in the codebase.")
    add_text(doc, "Containerization addresses one part of this gap by packaging an application with its runtime dependencies. Docker's container model has been widely used to reduce differences between developer machines and deployment hosts [3]. However, a Dockerfile cannot be generated safely from a file extension alone. A front-end project may need a build stage and a static server, while an API may need a process manager, a database connection, and a health endpoint. SentinelOps treats the container template as a reviewed artifact rather than a final answer.")
    add_text(doc, "Observability provides the evidence needed after an application starts. Prometheus stores metrics as labeled time-series data and collects them from monitored targets [4]. Grafana can group queries into dashboards that help a user inspect the service state [6]. In a Kubernetes environment, an HPA adjusts replicas against observed resource or custom metrics [5]. These tools are useful only when their assumptions are visible. A dashboard with no reachable target and an autoscaler with no reliable metrics do not make a service safer.")
    add_text(doc, "Recent work on AI-driven DevOps focuses on decision support, anomaly detection, and automated remediation [1]. SentinelOps adopts the decision-support part of that direction. It uses a recommendation layer to suggest a stack or deployment profile, while the repository model, generated files, commands, and warnings remain inspectable. This makes the desktop workflow appropriate for project work where users must explain their design choices during demonstrations, reviews, or submissions.")
    add_text(doc, "The project therefore addresses a practical question: how can a desktop application take an unknown Git repository through a transparent path from inspection to locally validated delivery assets? The answer is not a single model or a single deployment target. It is a sequence of evidence-based modules with clear handoffs and a place for a user to correct uncertain information.")

    add_heading(doc, "III DESIGN GOALS AND SCOPE")
    add_text(doc, "SentinelOps has five design goals. First, it must recognize common project shapes without assuming that every repository is a single service. Second, it must separate local verification from deployment generation so a broken application is visible early. Third, it must produce ordinary, editable artifacts rather than hiding configuration behind a graphical interface. Fourth, it must keep local machine changes limited: Docker Desktop is checked as an external prerequisite on Windows and is not installed by the application. Finally, it must make operational choices visible through a deployment profile, health checks, scaling settings, and an observability plan.")
    add_text(doc, "The prototype targets Node.js, React, Next.js, Vue, Angular, Python, Flask, FastAPI, Django, Go, and Rust projects. Its detector looks for package manifests, lock files, framework dependencies, source layout, scripts, ports, and common container configuration. The output is a proposed model, not a guarantee. Users can edit the detected language, framework, service type, runtime commands, ports, and selected CI provider before proceeding.")

    add_heading(doc, "IV SYSTEM ARCHITECTURE")
    add_text(doc, "Figure 1 summarizes the controlled flow. Repository intake produces a local workspace and Git metadata. The analysis module builds an evidence-backed project model. Local verification then installs declared dependencies, starts the service using the selected command, probes a health endpoint when one is available, and records URLs that a user can open in a browser. Only after that gate does the asset generator prepare container, pipeline, deployment, and monitoring files.")
    add_architecture_flow(doc)
    add_text(doc, "The orchestration layer does not replace a production platform. It coordinates choices made by the user. For example, a local profile can favor a native process for early testing, while a container profile can generate Compose assets for a reproducible local environment. A Kubernetes profile adds a Deployment, Service, and HorizontalPodAutoscaler manifest. Kubernetes describes the HPA as a controller that adjusts workload replicas according to observed metrics [5], so SentinelOps treats scaling thresholds as policy that must be reviewed, not as an unexplained automatic action.")
    add_text(doc, "Each project follows a state model. The initial state contains a repository address and a project name. After cloning, the repository model contains metadata such as the remote URL, current branch, commit identifier, file inventory, and detected components. The review state contains user-selected commands, ports, and a target profile. The validated state adds command output, a health result, and local URLs. The generated state records the paths of Docker, pipeline, deployment, and monitoring files. Storing these transitions gives the user a concise history of how a delivery decision was reached.")
    add_text(doc, "The desktop interface is organized around these states instead of forcing users through a hidden linear wizard. A repository can be re-analyzed after a correction, and a generated plan can be changed without recloning the project. This is important when a repository uses a nonstandard script name, a custom port, or a monorepo layout. The user should be able to amend the model at the point where the assumption is wrong rather than edit a downstream file and lose the connection to the original evidence.")
    add_text(doc, "The architecture also separates private information from generated configuration. Source URLs and local paths may be stored in the project record, while registry credentials, deployment secrets, and remote cluster tokens are never guessed by the detector. If a later integration needs these values, it should request them through the user interface and pass them through a secure host-specific storage mechanism. This separation reduces the chance that generated examples encourage users to commit secrets to a repository.")

    add_heading(doc, "V IMPLEMENTATION DETAILS")
    add_text(doc, "The desktop application is built around a narrow set of modules. The Git module clones a repository into a project-specific directory and supports deletion of that cloned workspace after confirmation. The detector reads repository files and calculates findings from observable evidence. It flags client and server directories separately when it finds independent package manifests or framework signatures. This allows the interface to display two start commands and two local addresses when a project needs them.")
    add_text(doc, "The validation module resolves the selected install command, starts the application with the selected runtime command, and captures output for the user. It reports the process state and the local URL instead of claiming that the application is ready solely because a command was launched. A successful probe is useful evidence, but it is still scoped to the chosen endpoint and local environment. This distinction prevents generated deployment assets from being confused with a verified production release.")
    add_text(doc, "For containerized execution, SentinelOps derives a Dockerfile template from the confirmed runtime. Docker packages an application with the libraries and configuration it needs, which helps reduce environment differences across development and deployment [3]. The generated file remains editable because build steps, secrets, multi-stage builds, and operating-system packages are project-specific. Docker image pulls and container runs are available only after the host reports that Docker is installed and usable.")
    add_text(doc, "The pipeline module selects GitHub Actions, GitLab CI, or Jenkins based on repository context and user choice. It generates a visible sequence for checkout, dependency installation, build, test, image build, and optional publish. Credentials, registry URLs, and production deployment commands are intentionally not inferred. They require values supplied by the user or organization.")
    add_text(doc, "Repository analysis begins with files that carry explicit project intent. For JavaScript and TypeScript projects, the detector examines package.json, package-manager lock files, declared scripts, and framework dependencies. For Python projects, it examines requirements files, pyproject metadata, common entry files, and framework imports. Comparable rules identify Go modules, Rust manifests, and existing container files. The detector then combines file evidence with directory evidence. A client folder with its own manifest and a server folder with a different manifest is reported as a split application rather than as one ambiguous project.")
    add_text(doc, "The recommended commands are derived from the inspected manifests. For example, a declared development script has stronger evidence than a generic framework default. SentinelOps records the source of a command in the project model, so the interface can distinguish a direct repository script from a fallback suggestion. This is useful when the user changes a command manually. The revised command becomes the selected value for local validation and later asset generation, while the original detection remains visible for comparison.")
    add_text(doc, "Local validation is intentionally separate from Docker execution. The application first checks whether the selected runtime is available, runs the dependency installation command in the cloned workspace, and starts the service with a managed process. Output is streamed into the desktop interface and is retained with the validation record. Where a web address can be inferred, SentinelOps probes the endpoint and displays a clickable local URL. For a project with separate client and server services, the result includes independent process states and URLs instead of a single generic success message.")
    add_text(doc, "The Docker service uses a host status check before it offers image-related actions. On Windows, SentinelOps tells the user that Docker Desktop must be installed and started; it does not attempt to install it or modify the host. When Docker is available, the project can pull a selected base image, build an application image, and run a container with the proposed port mapping. The generated Dockerfile and Compose file are shown as normal text assets so that the user can change base images, build arguments, environment variables, volumes, or service names before execution.")
    add_text(doc, "A free language model can be used as a secondary recommender for a stack, delivery path, or explanation of an error. The recommendation layer is bounded by the repository model. It should not invent a framework that the detector did not observe, execute shell commands without approval, or overwrite generated files without preserving the previous version. In this design, model output is treated as a proposal accompanied by evidence and confidence notes. Deterministic checks remain responsible for file discovery, command execution, and validation status.")
    add_text(doc, "The project includes a visualization view that presents the chosen technology stack and the planned delivery path. It maps repository findings to the selected CI provider, container base image, deployment profile, monitoring assets, and response policy. The purpose is operational clarity. A user can see, for example, that a FastAPI service was paired with Python, Docker, a GitHub Actions workflow, a Kubernetes profile, Prometheus scraping, and a CPU-based HPA before any remote deployment command is considered.")
    add_text(doc, "Error handling is designed around recovery rather than silent retry. A clone error keeps the project in the intake state and surfaces the Git message. A detection error preserves the repository record and asks the user to select fields manually. A failed install or local start stores the terminal output and prevents the interface from presenting deployment generation as a verified path. A Docker error reports the missing host prerequisite or engine result. These boundaries keep a failure trace close to the action that caused it.")
    add_module_table(doc)

    add_heading(doc, "VI USER WORKFLOW AND CONTROLLED EXECUTION")
    add_text(doc, "A user begins by creating a SentinelOps project and supplying a Git repository address. The application records a human-readable project name and chooses a local directory for the clone. If the repository already exists in a previous project workspace, the user can inspect its metadata or remove the cloned copy through a dedicated delete action. Deletion applies only to the copied workspace managed by SentinelOps; it must not remove a user's original repository or any remote repository. This small constraint matters because users often experiment with several repositories while learning the tool.")
    add_text(doc, "After cloning, the analysis screen presents the detected stack as editable fields rather than as a black-box result. It can list a primary language, framework, package manager, build tool, service type, detected ports, container support, and CI configuration. When the repository contains both a front end and a back end, the screen shows each component separately. A user can confirm a React client on one port and a Python API on another, then keep those selections when creating Docker or pipeline assets.")
    add_text(doc, "The next step is local validation. The user reviews the install and start commands before running them. SentinelOps then displays terminal output, process status, and local addresses. This workflow is especially useful for a project whose front end and back end must be launched separately. Instead of asking the user to infer the address from logs, the application shows an explicit link such as a client URL and an API URL when they are known. If a process stops unexpectedly, its exit result becomes part of the validation record.")
    add_text(doc, "Once validation succeeds or the user elects to continue with a known limitation, the planning screen recommends a delivery profile. The recommendation considers the selected stack, whether the service is split, whether a Dockerfile already exists, and the user's stated target. The screen explains its recommendation in ordinary language. For example, it may suggest native local execution for an early single-service prototype, Compose for a service with a local dependency, or Kubernetes only after a container image and cluster target have been confirmed.")
    add_text(doc, "The execution screen retains a visible distinction between generated files and commands that have actually run. Creating a Dockerfile does not mean an image exists. Generating a Jenkinsfile does not mean a pipeline has been configured on a Jenkins server. Preparing a Prometheus configuration does not mean metrics are being collected. This distinction prevents an attractive visualization from overstating system state. Every stage can display a generated, configured, started, validated, failed, or not-applicable status.")
    add_text(doc, "The final project view acts as a compact record for demonstrations. It can show the repository source, selected stack, local test result, container choice, CI/CD template, deployment profile, observability configuration, scaling policy, and recommended maintenance actions. The same record gives a reviewer enough context to understand why the application chose its assets. It also creates a useful handoff note if another developer needs to continue the setup without repeating the initial repository investigation.")

    add_heading(doc, "VII DELIVERY PROFILES AND CI CD")
    add_text(doc, "SentinelOps presents deployment as a choice among profiles rather than a one-size-fits-all action. The native local profile is the least complex option. It starts the selected command directly on the workstation and is useful for debugging a newly cloned repository. The Docker Compose profile packages the application and supporting services in a repeatable local bundle. The Kubernetes profile is intended for users who already have an approved cluster and understand the operational implications of a remote deployment. Table III summarizes the distinction.")
    add_profile_table(doc)
    add_text(doc, "The recommended profile depends on the confirmed project model. A small static front end can often be tested locally and then built into a compact web-serving image. A service that requires a database may benefit from Compose during development because the app and its dependency can be described together. A split front end and back end needs separate service entries, port mappings, and health checks. Kubernetes becomes a reasonable recommendation only when a cluster, image registry, and operational ownership are already available. SentinelOps states these prerequisites in the profile description instead of presenting Kubernetes as a default upgrade.")
    add_text(doc, "The CI/CD templates have the same principle. A GitHub repository can begin with GitHub Actions, a GitLab repository with GitLab CI, and a Jenkins-based environment with a Jenkinsfile. Each template starts with checkout and dependency installation, then runs the project build and selected tests. Container build and push steps remain optional because an image registry requires a user-supplied destination and credentials. The template can also generate a deployment stage, but it remains disabled until the user supplies the target environment and approval process.")
    add_text(doc, "The generated pipelines are designed to be read by a developer who did not build SentinelOps. Steps use clear names and standard repository scripts where possible. A pipeline should fail when a build or test fails, and the report should point back to the command output. The desktop application does not promise to repair every failure automatically. It can identify a failed stage, collect the message, explain common causes, and propose a next check. This is safer than retrying a build indefinitely or modifying source code without review.")
    add_text(doc, "Image tagging is treated as a release decision. A generated configuration can use a local tag for testing and can reserve immutable tags based on a commit identifier for an approved registry workflow. A later production integration should add provenance, vulnerability scanning, and access control before an image is promoted. SentinelOps can prepare the surrounding files, but the project owner remains responsible for registry permissions and the policy that authorizes an image to reach an environment.")

    add_heading(doc, "VIII OBSERVABILITY AND MAINTENANCE")
    add_text(doc, "Prometheus collects time-series metrics from configured targets, while Grafana provides dashboards that make those metrics easier to inspect [4], [6]. SentinelOps creates a Prometheus scrape configuration, a starter Grafana dashboard definition, and an optional Compose stack for local observation. The desktop interface can summarize generated endpoints and recommended signals such as service availability, CPU use, memory use, request errors, and restart count. It must not claim live metrics until the user has started the monitoring stack and the application exposes a compatible metrics endpoint.")
    add_text(doc, "The maintenance model is recommendation-led. When a health probe fails or an observed threshold is exceeded, the interface can propose a restart, rollback, scaling change, configuration review, or ticket. A runbook explains the suggested action and the evidence that triggered it. Automatic source-code repair and unapproved production changes are outside the prototype scope. This boundary makes an automated workflow more suitable for a desktop tool used by learners and small teams, where an incorrect assumption can otherwise change a local project or remote environment unexpectedly.")
    add_text(doc, "The observability configuration starts with a small, explicit target list. The user selects the application endpoint and the scrape interval, and SentinelOps writes a Prometheus configuration that can be reviewed before it is started. A starter Grafana dashboard groups availability, CPU, memory, process restarts, and request-error panels when the chosen runtime exposes those metrics. The dashboard definition is a starting point, not proof that every application has the same metric names. A Python API, a Node.js process, and a static client expose different signals and may require an instrumentation change before the panels show data.")
    add_text(doc, "Autoscaling is available only for the Kubernetes profile. SentinelOps generates an autoscaling v2 resource with reviewed minimum and maximum replica values and a target CPU utilization. It also includes resource requests in the generated workload because CPU-based scaling relies on those requests. The user can adjust the capacity policy in the interface and see the resulting manifest. The system explains that high traffic alone does not guarantee a scale-out event; the configured metric, target, workload readiness, and cluster capacity all affect the result.")
    add_text(doc, "A controlled self-healing loop uses four stages. First, a health or metric signal identifies an abnormal condition. Second, the system gathers context such as the selected profile, recent command output, configured endpoint, and the last known status. Third, it matches the event to a runbook suggestion. Finally, it requests user approval for any action that changes a running environment. This can include a local restart, a rollback to a previously selected configuration, a scaling adjustment, or a notification. The loop is useful because it makes the suggested response explainable and reversible.")
    add_text(doc, "Notifications are an integration boundary rather than an assumed service. The project model can reserve hooks for email, Slack, Microsoft Teams, Jira, or PagerDuty. A real connection requires credentials, endpoint configuration, and an organization-specific escalation policy. In the current project, the user-facing value is the event record and recommended action. This keeps the paper aligned with what SentinelOps can safely generate and avoids claiming that an alert was delivered when no integration has been configured.")
    add_text(doc, "The desktop visualization supports the same maintenance model. A generated plan can show the container image, CI/CD stages, deployment profile, metric sources, autoscaling range, and recent health result as connected modules. The visual view does not replace Grafana. Instead, it explains how the generated assets relate before the user opens a monitoring dashboard. That distinction is helpful during project demonstrations because viewers can follow the path from repository analysis to an operational policy without needing live cloud credentials.")

    add_heading(doc, "IX EVALUATION PLAN")
    add_text(doc, "The prototype should be evaluated with a curated set of public repositories that cover single-service Node.js and Python applications, front-end applications, and repositories with separate client and server folders. For each repository, an evaluator records the expected language, framework, package manager, start command, ports, and service boundaries before running SentinelOps. The primary measures are detection precision, detection recall, the rate of correct local launch instructions, and the number of generated artifacts that require manual correction.")
    add_text(doc, "A second evaluation should examine operational transparency. Participants should be able to identify the selected deployment profile, the command used for local validation, the container image tag, the metrics endpoint, and the condition that would request scaling. A short user study can compare this workflow with a manual checklist. The purpose is to find missing evidence and confusing defaults, not to make unsupported claims about deployment speed or reliability.")
    add_validation_table(doc)
    add_text(doc, "The evaluation can also score the quality of manual correction. An evaluator deliberately changes a command, a port, or a framework field and then checks whether the regenerated Dockerfile, pipeline, and deployment artifacts reflect the selected value. A good result is not merely a high automatic detection rate. It is a tool that remains coherent after a user corrects the detector. This matters because repositories often contain legacy scripts or examples that should not be chosen as a production entry point.")
    add_text(doc, "Failure-oriented tests are equally important. The test set should include an invalid Git URL, a private repository without access, a dependency-install failure, a missing runtime, a server that starts but does not answer the health probe, a missing Docker engine, and a missing metrics endpoint. For each case, the evaluator records whether SentinelOps preserves the project state, exposes the error source, avoids a misleading ready status, and offers a useful next action. These tests assess reliability of the workflow itself rather than performance of a particular application.")
    add_text(doc, "The paper does not report numeric outcomes because a full benchmark set and participant study have not yet been completed. That restraint is important. The prototype demonstrates the architecture and generated artifacts, while the proposed study defines how its accuracy and usability can later be measured. Results should include the repository sample, exact versions, selected manual overrides, and failures so that another student team can reproduce the comparison.")

    add_heading(doc, "X SECURITY LIMITATIONS AND FUTURE WORK")
    add_text(doc, "Repository inspection cannot reliably discover every runtime requirement. Private registries, undocumented environment variables, external databases, and shell scripts may require manual input. Dynamic stack selection also depends on the quality of the underlying evidence and any optional language model recommendation. SentinelOps therefore records both its detection and the user’s final selections. Future work can add signed artifact provenance, secret scanning, configurable policy checks, richer service discovery, and a controlled approval workflow for remote deployment.")

    add_text(doc, "Security requires further work before SentinelOps is used outside a controlled development setting. Generated containers should run with least privilege, use small maintained base images, and avoid embedding secrets. CI/CD credentials should be stored in the provider's secret mechanism rather than copied into generated YAML files. A production deployment path should include access review, image scanning, signed releases, and a rollback plan. These controls are intentionally described as requirements instead of being presented as completed capabilities.")
    add_text(doc, "The current scope also excludes autonomous code changes. A language model can help interpret logs or describe a likely configuration issue, but it may produce a plausible yet incorrect fix. SentinelOps therefore limits automated behavior to inspection, generation, local process control, health observation, and recommendations that a user can review. Future versions could add a sandboxed repair branch, a test-only validation gate, and a pull-request workflow for proposed changes. That approach would retain human review while making recovery experiments more systematic.")
    add_text(doc, "Another future direction is richer architecture discovery. A repository may contain infrastructure modules, database migrations, message queues, mobile clients, or multiple deployable services. Extending the project model to represent those components explicitly would improve deployment recommendations. The same model could support cost-aware choices, environment-specific configuration, and dependency graphs. Such enhancements should continue to preserve the current principle that generated conclusions are traceable to files, user selections, or stated operational policies.")

    add_heading(doc, "XI DEMONSTRATION SCENARIO")
    add_text(doc, "A representative SentinelOps session begins with a repository that contains a React client and a Python API. The user creates a project, pastes the repository URL, and starts the clone. The Git module displays the remote address and local workspace path after the operation completes. The detector then finds two independent package descriptions and reports a client component and a server component. It proposes a JavaScript front-end framework, a Python API framework, separate dependency commands, and two ports. The user can revise any of these values before they become part of the project plan.")
    add_text(doc, "The user next runs local validation. SentinelOps installs the front-end and back-end dependencies in their respective folders, starts each service, and streams the output into the workspace view. The result reports the front-end URL and the API URL separately. If the front end can load while an API route fails, the record shows that difference instead of treating the project as fully healthy. A developer can use the captured output to change an environment variable, port, or start command and then repeat the test. This scenario demonstrates why separate service discovery is useful for full-stack repositories.")
    add_text(doc, "After a successful local check, the planner can recommend Docker Compose as the initial deployment profile. It prepares one service definition for the client and one for the API, with their chosen ports and a shared network. The user reviews the proposed Dockerfiles, chooses a supported base image, and starts image pull or build actions only if Docker is ready on the host. The Docker status view makes the prerequisite visible on Windows. It does not offer a hidden installation path or claim that a container was built before the engine returns a result.")
    add_text(doc, "For source delivery, the user selects the CI provider that matches the repository host. A GitHub repository can receive a GitHub Actions workflow with checkout, dependency installation, build, test, and optional image build stages. A GitLab repository can receive a comparable GitLab CI definition. A Jenkins environment can receive a Jenkinsfile that its administrator may import into an existing server. The template remains a generated file in the project workspace. Connecting it to a remote service is a later deliberate action that needs credentials and repository permissions.")
    add_text(doc, "If the user later chooses a Kubernetes profile, SentinelOps converts the selected service information into editable workload manifests. The plan can include a Deployment, Service, resource requests, and an HPA with an explicit replica range. The interface explains that this profile requires a cluster, an image registry, and access approval. It also makes clear that the HPA uses the selected resource policy. It does not imply that traffic automatically produces a cloud deployment or that the desktop app can scale a cluster it cannot access.")
    add_text(doc, "The observability step adds a local Prometheus configuration and a starter Grafana dashboard definition. Once the user starts the monitoring services and exposes application metrics, the plan can connect the application target to dashboard panels. A failed scrape is shown as a missing dependency or configuration issue, not as a successful observation. The maintenance view can then associate a failed health check or high utilization signal with a runbook entry. The user sees the evidence, suggested action, and whether approval is required before any corrective operation.")
    add_text(doc, "This scenario illustrates the intended meaning of automation in SentinelOps. The desktop application removes repetitive discovery and configuration work, but it does not hide decisions that carry operational risk. A student can demonstrate the entire flow from repository to generated artifacts on one machine. A small team can use the same flow to create a consistent starting point before involving a registry, CI server, cluster administrator, or production monitoring environment.")

    add_heading(doc, "XII REPRODUCIBILITY AND MAINTAINABILITY")
    add_text(doc, "The project favors text-based outputs because they are easy to inspect, diff, and version with the source code. Dockerfiles, Compose files, CI pipelines, Kubernetes manifests, Prometheus configurations, Grafana dashboard JSON, and runbooks should be stored in a predictable project folder. SentinelOps records the source repository revision used for analysis and the selected settings used for generation. If a user regenerates an asset after changing the stack model, the application can compare the new configuration with the earlier version and show which assumptions changed.")
    add_text(doc, "A maintainable tool also needs predictable cleanup. SentinelOps gives each cloned repository its own project workspace and provides a specific option to delete that cloned copy. It should stop managed local processes before removing the workspace and should clearly name the target directory before confirmation. Generated assets can be kept even when the clone is removed if the user wants to retain a design record. This behavior lets users recover disk space without losing the history of their chosen deployment plan.")
    add_text(doc, "For academic use, repeatability depends on documenting the host assumptions. A project demonstration should state the operating system, runtime versions, Docker availability, repository commit, selected commands, and chosen profile. Screenshots or exported project summaries can capture the detected stack and generated flow. These details make it possible for another evaluator to distinguish a product limitation from a missing host prerequisite, and they make the proposed evaluation in Section IX more credible.")

    add_heading(doc, "XIII CONCLUSION")
    add_text(doc, "SentinelOps provides a practical path from a Git repository to a reviewed DevOps plan. Its central idea is simple: inspect the project, prove that it can run locally, then generate transparent assets for containers, CI/CD, deployment, and monitoring. The prototype keeps users in control of uncertain decisions and clearly separates generated files from verified production operations. That combination makes the system appropriate for learning, project demonstrations, and small software teams that need a repeatable starting point for delivery automation.")

    add_heading(doc, "REFERENCES")
    add_references(doc)

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_document()
def _font(size, bold=False):
    name = "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"
    return ImageFont.truetype(name, size) if Path(name).exists() else ImageFont.load_default()


def build_architecture_figure():
    """Create an original, source-derived diagram of SentinelOps modules."""
    image = Image.new("RGB", (1600, 720), "white")
    draw = ImageDraw.Draw(image)
    title_font, body_font, small_font = _font(30, True), _font(22), _font(18)
    draw.text((800, 24), "SENTINELOPS REVIEWED DELIVERY WORKFLOW", fill="#17324D", font=title_font, anchor="ma")
    boxes = [
        (55, 225, 330, 440, "Repository intake", "Clone Git source | Record branch and commit | Create isolated workspace", "#E7F1FA"),
        (405, 225, 680, 440, "Evidence analysis", "Detect stack and services | Infer commands and ports | Allow user corrections", "#E8F5F1"),
        (755, 225, 1030, 440, "Local validation", "Install dependencies | Run selected service | Probe health and URLs", "#FFF4DF"),
        (1105, 140, 1510, 315, "Generated assets", "Docker and Compose | CI/CD workflow | Kubernetes manifests", "#F1ECFA"),
        (1105, 405, 1510, 580, "Operations plan", "Prometheus and Grafana | HPA policy and runbook | User-approved actions", "#FCEAF0"),
    ]
    for x1, y1, x2, y2, heading, body, fill in boxes:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=22, fill=fill, outline="#315B7D", width=3)
        draw.text(((x1+x2)//2, y1+34), heading, fill="#17324D", font=body_font, anchor="ma")
        for index, line in enumerate(body.split(" | ")):
            draw.text(((x1+x2)//2, y1+84+index*35), line, fill="#283746", font=small_font, anchor="ma")
    def arrow(start, end):
        draw.line((start, end), fill="#315B7D", width=7)
        draw.polygon([(end[0], end[1]), (end[0]-18, end[1]-12), (end[0]-18, end[1]+12)], fill="#315B7D")
    arrow((330, 332), (400, 332)); arrow((680, 332), (750, 332))
    arrow((1030, 285), (1097, 227)); arrow((1030, 385), (1097, 492))
    draw.text((800, 640), "Generated configuration is reviewable; remote deployment and remediation remain user-controlled.", fill="#4B5F70", font=small_font, anchor="ma")
    image.save(FIGURE_PATH, "PNG")


